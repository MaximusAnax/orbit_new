"""Metric implementations for the datasweep eval suite (EVALS.md §3).

Ground truth never comes from the system under test:

* detection truth is the corrupter's manifest,
* repair truth is the **golden cell value**, read here with a plain stdlib
  parser rather than datasweep's reader,
* tier truth is the hand-authored ``traps/forbidden.json``,
* M5's truth is the parsed original itself and M6's truth is byte identity.

Every detection metric reads the persisted ``findings.jsonl`` — the artifact a
downstream consumer would actually get (EVALS.md §3 "Detection input") — never
``IssueSummary.samples`` and never engine internals.

The whole suite is hermetic: offline adapters, an in-memory repository,
``FixedClock("2026-01-01T00:00:00Z")``, committed fixtures, no network.
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Any

from datasweep.adapters.clock import FixedClock
from datasweep.adapters.notifier import NullNotifier
from datasweep.services import DatasweepService
from datasweep.store.memory import InMemoryRepository

FIXTURES = Path(__file__).resolve().parent / "fixtures"
GOLDEN_DIR = FIXTURES / "golden"
CORRUPT_DIR = FIXTURES / "corrupted"
TRAP_DIR = FIXTURES / "traps"

EVAL_CLOCK = "2026-01-01T00:00:00Z"

#: Manifest op → issue class (EVALS.md §3).  Fixed here, not derived.
OP_CLASS: dict[str, str] = {
    "mojibake_encode": "ENC",
    "pad_whitespace": "WS",
    "insert_nbsp": "WS",
    "insert_zero_width": "WS",
    "double_internal_space": "WS",
    "sentinel_missing": "MISS",
    "thousands_sep": "TYPE",
    "decimal_comma": "TYPE",
    "currency_prefix": "TYPE",
    "leading_apostrophe": "TYPE",
    "date_reformat": "DATE",
    "excel_serial": "DATE",
    "case_label": "CAT",
    "punct_label": "CAT",
    "typo_label": "CAT",
    "duplicate_row": "DUP",
    "inject_outlier": "OUT",
}

#: The eight detector classes M2 scores.  ``STR`` is reader-emitted, not a
#: detector family (SCOPE.md FR-6), so it is excluded by design.
M2_CLASSES: tuple[str, ...] = ("ENC", "WS", "MISS", "TYPE", "DATE", "CAT", "DUP", "OUT")

#: Report-only, class OUT by design (SCOPE.md D11) — excluded from M4's
#: fixable set and scored by M2 detection instead.
UNFIXABLE_CLASSES: frozenset[str] = frozenset({"OUT"})

#: The five classes whose ops are *entirely* auto-expected under the pinned
#: mix, so demoting any one of them is caught even if the aggregate survives.
REPAIR_MIN_CLASSES: tuple[str, ...] = ("ENC", "WS", "MISS", "TYPE", "DUP")

#: Column-scope "I cannot decide" rules (EVALS.md §3).
AMBIGUITY_RULES: frozenset[str] = frozenset(
    {"fix.date_canon_ambiguous", "fix.number_canon_ambiguous"}
)

#: Column-scope advisories: true statements about a column rather than claims
#: about individual cells, so M2 ignores them entirely (EVALS.md §3).
ADVISORY_RULES: frozenset[str] = frozenset(
    {"detect.mixed_number_conventions", "fix.currency_mixed"}
)


# --------------------------------------------------------------------------
# Gates (EVALS.md §5)
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Gate:
    key: str
    label: str
    threshold: float
    comparison: str  # ">=" or "=="
    rationale: str

    def passes(self, value: float) -> bool:
        return value >= self.threshold if self.comparison == ">=" else value == self.threshold


GATES: dict[str, Gate] = {
    gate.key: gate
    for gate in (
        Gate("M1", "M1 type_inference_accuracy", 0.95, ">=", "91 labeled column instances"),
        Gate("M2", "M2 detection macro-F1", 0.85, ">=", "recall and precision on all 8 classes"),
        Gate("M2_min", "M2_min per-class F1", 0.70, ">=", "no dead detector behind the mean"),
        Gate("M3", "M3 auto_fix_precision", 0.98, ">=", "the safety contract (SCOPE.md D1)"),
        Gate("M3_trap", "M3_trap trap-cell auto changes", 0.0, "==", "hand-vetted ambiguity"),
        Gate("M3_clean", "M3_clean clean-file auto changes", 0.0, "==", "cleaning is idempotent"),
        Gate(
            "M3_clean_findings",
            "M3_clean_findings clean-file findings + review items",
            0.0,
            "==",
            "the false-alarm half of H1",
        ),
        Gate("M4_auto", "M4 repair_auto", 0.85, ">=", "blocks broad timidity"),
        Gate("M4_auto_min", "M4 repair_auto_min", 0.85, ">=", "blocks demoting one class"),
        Gate("M4_total", "M4 repair_total", 0.90, ">=", "the review tier carries real proposals"),
        Gate("M5", "M5 reversibility", 1.0, "==", "the non-destructive contract is constructive"),
        Gate("M6", "M6 determinism", 1.0, "==", "cross-process, two PYTHONHASHSEED values"),
        Gate("M7", "M7 trap_disposition", 1.0, "==", "the positive half of the tier contract"),
    )
}


@dataclass(frozen=True)
class MetricResult:
    name: str
    value: float
    gate: float | None
    comparison: str
    passed: bool
    detail: dict[str, Any] = field(default_factory=dict)


# --------------------------------------------------------------------------
# Fixture loading — plain stdlib parsers, independent of datasweep's readers
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class Table:
    headers: list[str]
    rows: list[list[str | None]]

    def cell(self, row: int, col: int) -> str | None:
        if 0 <= row < len(self.rows) and 0 <= col < len(self.rows[row]):
            return self.rows[row][col]
        return None


def read_table(path: Path) -> Table:
    """Read a fixture with the stdlib, mirroring ``generate.py``'s writers.

    Deliberately not datasweep's reader stack: repair truth must not be
    mediated by the code under test.
    """
    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        headers: list[str] = []
        records: list[dict[str, Any]] = []
        for line in path.read_text(encoding="utf-8").splitlines():
            if not line.strip():
                continue
            record = json.loads(line, parse_float=str, parse_int=str)
            records.append(record)
            for key in record:
                if key not in headers:
                    headers.append(key)
        rows = [
            [None if record.get(key) is None else str(record.get(key)) for key in headers]
            for record in records
        ]
        return Table(headers=headers, rows=rows)

    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            grid = [
                [None if cell is None else str(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
        finally:
            workbook.close()
        while grid and all(cell is None for cell in grid[-1]):
            grid.pop()
        width = len(grid[0]) if grid else 0
        return Table(
            headers=[cell or "" for cell in grid[0][:width]],
            rows=[list(row[:width]) for row in grid[1:]],
        )

    import csv

    delimiter = "\t" if suffix == ".tsv" else ","
    text = path.read_text(encoding="utf-8")
    reader = csv.reader(text.splitlines(), delimiter=delimiter, quotechar='"')
    grid = [row for row in reader if row]
    return Table(
        headers=list(grid[0]),
        rows=[[None if cell == "" else cell for cell in row] for row in grid[1:]],
    )


def corrupted_fixtures() -> list[Path]:
    return sorted(p for p in CORRUPT_DIR.iterdir() if p.suffix != ".json")


def golden_fixtures() -> list[Path]:
    return sorted(GOLDEN_DIR.iterdir())


def trap_fixtures() -> list[Path]:
    return sorted(p for p in TRAP_DIR.glob("*.csv"))


def manifest_for(path: Path) -> dict[str, Any]:
    stem = path.name.rpartition(".")[0]
    return json.loads((CORRUPT_DIR / f"{stem}.manifest.json").read_text(encoding="utf-8"))


@lru_cache(maxsize=1)
def type_labels() -> dict[str, dict[str, str]]:
    payload = json.loads((FIXTURES / "types.json").read_text(encoding="utf-8"))
    return {key: value for key, value in payload.items() if not key.startswith("_")}


@lru_cache(maxsize=1)
def forbidden() -> dict[str, list[dict[str, Any]]]:
    payload = json.loads((TRAP_DIR / "forbidden.json").read_text(encoding="utf-8"))
    return {key: value for key, value in payload.items() if not key.startswith("_")}


@lru_cache(maxsize=1)
def expected() -> dict[str, Any]:
    return json.loads((FIXTURES / "expected.json").read_text(encoding="utf-8"))


# --------------------------------------------------------------------------
# Running the pipeline over a fixture
# --------------------------------------------------------------------------


@dataclass
class Outcome:
    """Everything one fixture run produced, read back from its artifacts."""

    name: str
    kind: str  # corrupted | golden | trap
    path: Path
    findings: list[dict[str, Any]]
    audit: list[dict[str, Any]]
    review_items: list[dict[str, Any]]
    profiles: dict[str, str]
    issue_counts: dict[str, int]
    change_counts: dict[str, int]
    revert_ok: bool
    golden_table: Table | None = None
    manifest: dict[str, Any] | None = None

    @property
    def auto_changes(self) -> list[dict[str, Any]]:
        return [entry for entry in self.audit if entry.get("tier") == "auto"]

    def golden_cell(self, row: int, col: int) -> str | None:
        assert self.golden_table is not None
        return self.golden_table.cell(row, col)


def _service() -> DatasweepService:
    return DatasweepService(
        InMemoryRepository(),
        clock=FixedClock(EVAL_CLOCK, step_seconds=1),
        notifier=NullNotifier(),
    )


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    return [
        json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()
    ]


def run_fixture(path: Path, kind: str, workdir: Path) -> Outcome:
    """Clean one fixture with offline adapters and read its artifacts back."""
    service = _service()
    out = workdir / kind / path.stem
    run = service.clean_file(str(path), out=str(out))
    directory = Path(run.artifact_dir or "")
    audit = _read_jsonl(directory / "audit.jsonl")
    manifest = manifest_for(path) if kind == "corrupted" else None
    golden_table = (
        read_table(GOLDEN_DIR / manifest["golden"]) if manifest is not None else read_table(path)
    )
    return Outcome(
        name=path.name,
        kind=kind,
        path=path,
        findings=_read_jsonl(directory / "findings.jsonl"),
        audit=[entry for entry in audit if entry.get("kind") != "header"],
        review_items=[item.model_dump(mode="json") for item in service.review_items(run.id)],
        profiles={
            profile.name: profile.inferred_type.value
            for profile in service.repo.list_column_profiles(run.id)
        },
        issue_counts=dict(run.issue_counts),
        change_counts=dict(run.change_counts),
        revert_ok=service.revert_check(run.id).match,
        golden_table=golden_table,
        manifest=manifest,
    )


def revision_two_reversibility(workdir: Path) -> tuple[str, bool]:
    """M5's extra case: accept-all on a review-bearing fixture, then revert r2."""
    fixture = CORRUPT_DIR / "contacts.c1.csv"
    service = _service()
    run = service.clean_file(str(fixture), out=str(workdir / "revision"))
    items = [item.id for item in service.review_items(run.id)]
    assert items, "the revision-2 case needs a fixture with a non-empty review queue"
    revision = service.decide(run.id, accept=items)
    assert revision.revision_no == 2
    return fixture.name, service.revert_check(run.id, 2).match


# --------------------------------------------------------------------------
# Manifest helpers
# --------------------------------------------------------------------------


def injected_ops(outcome: Outcome) -> list[dict[str, Any]]:
    assert outcome.manifest is not None
    return list(outcome.manifest["ops"])


def op_class(op: dict[str, Any]) -> str:
    return OP_CLASS[op["op"]]


def column_truth(outcome: Outcome, col_name: str | None) -> dict[str, Any]:
    if outcome.manifest is None or col_name is None:
        return {}
    return outcome.manifest.get("column_truth", {}).get(col_name, {})


# --------------------------------------------------------------------------
# M1 — type inference accuracy
# --------------------------------------------------------------------------


def m1_type_inference(outcomes: list[Outcome]) -> MetricResult:
    labels = type_labels()
    correct = 0
    total = 0
    errors: list[str] = []
    for outcome in outcomes:
        if outcome.kind != "corrupted":
            continue
        assert outcome.manifest is not None
        expected_types = labels[outcome.manifest["golden"]]
        for column, label in expected_types.items():
            total += 1
            actual = outcome.profiles.get(column)
            if actual == label:
                correct += 1
            else:
                errors.append(f"{outcome.name}:{column} expected {label}, got {actual}")
    value = correct / total if total else 0.0
    gate = GATES["M1"]
    return MetricResult(
        name="M1_type_inference_accuracy",
        value=value,
        gate=gate.threshold,
        comparison=gate.comparison,
        passed=gate.passes(value),
        detail={"labeled_columns": total, "correct": correct, "errors": errors[:10]},
    )


# --------------------------------------------------------------------------
# M2 — detection F1 per class
# --------------------------------------------------------------------------


def _finding_key(finding: dict[str, Any]) -> tuple[int | None, int | None]:
    return (finding.get("row"), finding.get("col"))


def m2_detection(outcomes: list[Outcome]) -> tuple[MetricResult, MetricResult]:
    injected: dict[str, int] = Counter()
    detected: dict[str, int] = Counter()
    reports: dict[str, int] = Counter()
    matched: dict[str, int] = Counter()

    for outcome in outcomes:
        if outcome.kind == "trap":
            continue  # trap truth is forbidden.json (M3_trap / M7), not the manifest

        ops = injected_ops(outcome) if outcome.kind == "corrupted" else []
        by_cell: dict[tuple[str, int, int | None], dict[str, Any]] = {}
        by_row: dict[tuple[str, int], dict[str, Any]] = {}
        for op in ops:
            klass = op_class(op)
            injected[klass] += 1
            if klass == "DUP":
                by_row[(klass, op["row"])] = op
            else:
                by_cell[(klass, op["row"], op["col"])] = op
        covered: set[int] = set()

        # Column-scope ambiguity findings are one report unit per column.
        ambiguity_groups: dict[tuple[str, int | None], list[dict[str, Any]]] = defaultdict(list)
        for finding in outcome.findings:
            rule, klass = finding["rule"], finding["klass"]
            if rule in ADVISORY_RULES or klass not in M2_CLASSES:
                continue
            if rule in AMBIGUITY_RULES and finding.get("evidence", {}).get("scope") == "column":
                ambiguity_groups[(klass, finding.get("col"))].append(finding)
                continue

            reports[klass] += 1
            if klass == "DUP":
                op = by_row.get((klass, finding.get("row", -1)))
            else:
                op = by_cell.get((klass, finding.get("row", -1), finding.get("col")))
            if op is None:
                continue
            matched[klass] += 1
            marker = id(op)
            if marker not in covered:
                covered.add(marker)
                detected[klass] += 1

        for (klass, col), group in sorted(
            ambiguity_groups.items(), key=lambda item: (item[0][0], item[0][1] or -1)
        ):
            # EVALS.md §3: a column-scope "I cannot decide" finding covers the
            # column's ops only where the manifest marks the column ambiguous.
            # Every corrupted date and numeric column is provable by
            # construction, so declaring one ambiguous is a detection error and
            # is charged as exactly one class false positive — which is what
            # stops "call everything ambiguous" from earning perfect recall.
            reports[klass] += 1
            truth = column_truth(outcome, group[0].get("col_name"))
            if "ambiguous" not in {truth.get("date_order"), truth.get("number_convention")}:
                continue
            matched[klass] += 1
            for op in ops:
                if op_class(op) == klass and op.get("col") == col and id(op) not in covered:
                    covered.add(id(op))
                    detected[klass] += 1

    per_class: dict[str, dict[str, float]] = {}
    f1_values: list[float] = []
    for klass in M2_CLASSES:
        recall = detected[klass] / injected[klass] if injected[klass] else 0.0
        precision = matched[klass] / reports[klass] if reports[klass] else 0.0
        f1 = 2 * precision * recall / (precision + recall) if (precision + recall) else 0.0
        per_class[klass] = {
            "injected": injected[klass],
            "detected": detected[klass],
            "reports": reports[klass],
            "precision": round(precision, 6),
            "recall": round(recall, 6),
            "f1": round(f1, 6),
        }
        f1_values.append(f1)

    macro = sum(f1_values) / len(f1_values) if f1_values else 0.0
    minimum = min(f1_values) if f1_values else 0.0
    macro_gate, min_gate = GATES["M2"], GATES["M2_min"]
    return (
        MetricResult(
            name="M2_detection_macro_f1",
            value=macro,
            gate=macro_gate.threshold,
            comparison=macro_gate.comparison,
            passed=macro_gate.passes(macro),
            detail={"per_class": per_class},
        ),
        MetricResult(
            name="M2_min_class_f1",
            value=minimum,
            gate=min_gate.threshold,
            comparison=min_gate.comparison,
            passed=min_gate.passes(minimum),
            detail={"weakest": min(per_class, key=lambda k: per_class[k]["f1"])},
        ),
    )


# --------------------------------------------------------------------------
# M3 — auto-fix precision and the three companion hard checks
# --------------------------------------------------------------------------


def _auto_change_is_correct(outcome: Outcome, entry: dict[str, Any]) -> bool:
    kind = entry.get("kind")
    if kind == "cell_change":
        row, col = entry.get("row"), entry.get("col")
        if row is None or col is None:
            return False
        return entry.get("after") == outcome.golden_cell(row, col)
    if kind == "row_drop":
        assert outcome.manifest is not None
        duplicates = {op["row"] for op in outcome.manifest["ops"] if op["op"] == "duplicate_row"}
        return entry.get("row") in duplicates
    # row_pad / header_rename / column_add: the fixtures are rectangular with
    # unique headers, so any structural change is a change golden does not have.
    return False


def m3_auto_precision(outcomes: list[Outcome]) -> MetricResult:
    correct = 0
    total = 0
    wrong: list[str] = []
    for outcome in outcomes:
        if outcome.kind != "corrupted":
            continue
        for entry in outcome.auto_changes:
            total += 1
            if _auto_change_is_correct(outcome, entry):
                correct += 1
            else:
                wrong.append(
                    f"{outcome.name} r{entry.get('row')}c{entry.get('col')} "
                    f"{entry.get('rule')}: {entry.get('after')!r} != "
                    f"{outcome.golden_cell(entry.get('row', 0), entry.get('col', 0))!r}"
                )
    value = correct / total if total else 1.0
    gate = GATES["M3"]
    return MetricResult(
        name="M3_auto_fix_precision",
        value=value,
        gate=gate.threshold,
        comparison=gate.comparison,
        passed=gate.passes(value),
        detail={"auto_changes": total, "correct": correct, "wrong": wrong[:10]},
    )


def forbidden_cells(trap_name: str) -> set[tuple[int, int]]:
    cells: set[tuple[int, int]] = set()
    for expectation in forbidden().get(trap_name, []):
        if expectation.get("forbid_auto"):
            cells.update((row, col) for row, col in expectation["cells"])
    return cells


def m3_trap(outcomes: list[Outcome]) -> MetricResult:
    violations: list[str] = []
    for outcome in outcomes:
        if outcome.kind != "trap":
            continue
        banned = forbidden_cells(outcome.name)
        banned_rows = {row for row, _ in banned}
        for entry in outcome.auto_changes:
            row, col = entry.get("row"), entry.get("col")
            row_scope = entry.get("kind") in {"row_drop", "row_pad"}
            hit = (row, col) in banned or (row_scope and row in banned_rows)
            if hit:
                violations.append(f"{outcome.name} r{row}c{col} {entry.get('rule')}")
    value = float(len(violations))
    gate = GATES["M3_trap"]
    return MetricResult(
        name="M3_trap_auto_changes",
        value=value,
        gate=gate.threshold,
        comparison=gate.comparison,
        passed=gate.passes(value),
        detail={"violations": violations},
    )


def m3_clean(outcomes: list[Outcome]) -> MetricResult:
    changes = [
        f"{outcome.name} r{entry.get('row')}c{entry.get('col')} {entry.get('rule')}"
        for outcome in outcomes
        if outcome.kind == "golden"
        for entry in outcome.auto_changes
    ]
    value = float(len(changes))
    gate = GATES["M3_clean"]
    return MetricResult(
        name="M3_clean_auto_changes",
        value=value,
        gate=gate.threshold,
        comparison=gate.comparison,
        passed=gate.passes(value),
        detail={"changes": changes[:10]},
    )


def m3_clean_findings(outcomes: list[Outcome]) -> MetricResult:
    total = 0
    breakdown: dict[str, int] = {}
    for outcome in outcomes:
        if outcome.kind != "golden":
            continue
        count = len(outcome.findings) + len(outcome.review_items)
        breakdown[outcome.name] = count
        total += count
    value = float(total)
    gate = GATES["M3_clean_findings"]
    return MetricResult(
        name="M3_clean_findings",
        value=value,
        gate=gate.threshold,
        comparison=gate.comparison,
        passed=gate.passes(value),
        detail={"per_file": breakdown},
    )


# --------------------------------------------------------------------------
# M4 — repair rate
# --------------------------------------------------------------------------


def _auto_repairs(outcome: Outcome) -> tuple[set[tuple[int, int]], set[int]]:
    """(cells restored to golden by an auto change, rows dropped by an auto change)."""
    cells: set[tuple[int, int]] = set()
    rows: set[int] = set()
    for entry in outcome.auto_changes:
        if entry.get("kind") == "cell_change":
            row, col = entry.get("row"), entry.get("col")
            if row is None or col is None:
                continue
            if entry.get("after") == outcome.golden_cell(row, col):
                cells.add((row, col))
        elif entry.get("kind") == "row_drop" and entry.get("row") is not None:
            rows.add(int(entry["row"]))
    return cells, rows


def _review_repairs(outcome: Outcome) -> set[tuple[int, int]]:
    """Cells whose *recommended* review proposal would restore golden."""
    cells: set[tuple[int, int]] = set()
    for item in outcome.review_items:
        for cell in item.get("proposal", {}).get("cells", []):
            row, col = cell.get("row"), cell.get("col")
            if row is None or col is None:
                continue
            if cell.get("after") == outcome.golden_cell(row, col):
                cells.add((row, col))
    return cells


def m4_repair(outcomes: list[Outcome]) -> tuple[MetricResult, MetricResult, MetricResult]:
    fixable = 0
    auto_ok = 0
    total_ok = 0
    per_class_total: Counter[str] = Counter()
    per_class_auto: Counter[str] = Counter()
    per_class_any: Counter[str] = Counter()

    for outcome in outcomes:
        if outcome.kind != "corrupted":
            continue
        auto_cells, auto_rows = _auto_repairs(outcome)
        review_cells = _review_repairs(outcome)
        for op in injected_ops(outcome):
            klass = op_class(op)
            if klass in UNFIXABLE_CLASSES:
                continue
            fixable += 1
            per_class_total[klass] += 1
            if klass == "DUP":
                repaired_auto = op["row"] in auto_rows
                repaired_any = repaired_auto
            else:
                key = (op["row"], op["col"])
                repaired_auto = key in auto_cells
                repaired_any = repaired_auto or key in review_cells
            if repaired_auto:
                auto_ok += 1
                per_class_auto[klass] += 1
            if repaired_any:
                total_ok += 1
                per_class_any[klass] += 1

    auto_rate = auto_ok / fixable if fixable else 0.0
    total_rate = total_ok / fixable if fixable else 0.0
    per_class = {
        klass: round(per_class_auto[klass] / per_class_total[klass], 6)
        for klass in sorted(per_class_total)
    }
    min_classes = [per_class.get(klass, 0.0) for klass in REPAIR_MIN_CLASSES]
    auto_min = min(min_classes) if min_classes else 0.0

    gate_auto, gate_min, gate_total = GATES["M4_auto"], GATES["M4_auto_min"], GATES["M4_total"]
    return (
        MetricResult(
            name="M4_repair_auto",
            value=auto_rate,
            gate=gate_auto.threshold,
            comparison=gate_auto.comparison,
            passed=gate_auto.passes(auto_rate),
            detail={"fixable_ops": fixable, "repaired": auto_ok, "per_class": per_class},
        ),
        MetricResult(
            name="M4_repair_auto_min",
            value=auto_min,
            gate=gate_min.threshold,
            comparison=gate_min.comparison,
            passed=gate_min.passes(auto_min),
            detail={"classes": {k: per_class.get(k, 0.0) for k in REPAIR_MIN_CLASSES}},
        ),
        MetricResult(
            name="M4_repair_total",
            value=total_rate,
            gate=gate_total.threshold,
            comparison=gate_total.comparison,
            passed=gate_total.passes(total_rate),
            detail={
                "fixable_ops": fixable,
                "repaired": total_ok,
                "per_class": {
                    klass: round(per_class_any[klass] / per_class_total[klass], 6)
                    for klass in sorted(per_class_total)
                },
            },
        ),
    )


# --------------------------------------------------------------------------
# M5 — reversibility
# --------------------------------------------------------------------------


def m5_reversibility(outcomes: list[Outcome], revision: tuple[str, bool]) -> MetricResult:
    checks = [
        (outcome.name, outcome.revert_ok)
        for outcome in outcomes
        if outcome.kind in {"corrupted", "golden"}
    ]
    checks.append((f"{revision[0]} (revision 2)", revision[1]))
    passed = sum(1 for _, ok in checks if ok)
    value = passed / len(checks) if checks else 0.0
    gate = GATES["M5"]
    return MetricResult(
        name="M5_reversibility",
        value=value,
        gate=gate.threshold,
        comparison=gate.comparison,
        passed=gate.passes(value),
        detail={"runs": len(checks), "failures": [name for name, ok in checks if not ok]},
    )


# --------------------------------------------------------------------------
# M6 — determinism, across two processes with different PYTHONHASHSEED
# --------------------------------------------------------------------------

M6_FIXTURES: tuple[str, ...] = (
    "contacts.c1.csv",
    "survey.c1.tsv",
    "inventory.c1.xlsx",
    "sensors.c1.jsonl",
)

ARTIFACT_NAMES: tuple[str, ...] = ("audit.jsonl", "findings.jsonl", "report.md")


def _clean_in_subprocess(fixture: Path, out: Path, hash_seed: str) -> dict[str, Any]:
    """Invoke the CLI in a fresh interpreter with a pinned PYTHONHASHSEED.

    The subprocess split is load-bearing (EVALS.md M6): Python's string hash
    randomization varies only *between* processes, so two in-process runs would
    agree even if output order depended on set or dict iteration.
    """
    environment = dict(os.environ)
    environment["PYTHONHASHSEED"] = hash_seed
    environment["PYTHONIOENCODING"] = "utf-8"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "datasweep.cli",
            "clean",
            str(fixture),
            "--out",
            str(out),
            "--db",
            str(out / "datasweep.db"),
            "--json",
        ],
        capture_output=True,
        text=True,
        env=environment,
        check=False,
    )
    if result.returncode != 0:  # pragma: no cover - surfaced as a failing gate
        raise RuntimeError(f"determinism subprocess failed: {result.stderr[-2000:]}")
    return json.loads(result.stdout)


def _recount_from_artifacts(directory: Path) -> tuple[dict[str, int], dict[str, int]]:
    findings = _read_jsonl(directory / "findings.jsonl")
    audit = [e for e in _read_jsonl(directory / "audit.jsonl") if e.get("kind") != "header"]
    issue_counts = Counter(finding["klass"] for finding in findings)
    change_counts = {
        "auto": sum(1 for entry in audit if entry.get("tier") == "auto"),
        "review": sum(1 for finding in findings if finding["tier"] == "review"),
        "report": sum(1 for finding in findings if finding["tier"] == "report"),
    }
    return dict(issue_counts), change_counts


def m6_determinism(workdir: Path) -> MetricResult:
    problems: list[str] = []
    for name in M6_FIXTURES:
        fixture = CORRUPT_DIR / name
        left = workdir / "m6" / "seed0" / fixture.stem
        right = workdir / "m6" / "seed1" / fixture.stem
        run_left = _clean_in_subprocess(fixture, left, "0")
        run_right = _clean_in_subprocess(fixture, right, "1")
        dir_left, dir_right = Path(run_left["artifact_dir"]), Path(run_right["artifact_dir"])
        if dir_left.name != dir_right.name:
            problems.append(f"{name}: artifact directory names differ")
        cleaned = sorted(p.name for p in dir_left.glob("cleaned.*"))
        for artifact in [*cleaned, *ARTIFACT_NAMES]:
            left_bytes = (dir_left / artifact).read_bytes()
            right_bytes = (dir_right / artifact).read_bytes()
            if left_bytes != right_bytes:
                problems.append(f"{name}: {artifact} differs across PYTHONHASHSEED")

        issue_counts, change_counts = _recount_from_artifacts(dir_left)
        if issue_counts != {k: v for k, v in run_left["issue_counts"].items() if v}:
            problems.append(f"{name}: issue_counts do not match findings.jsonl")
        if change_counts != run_left["change_counts"]:
            problems.append(f"{name}: change_counts do not match the artifacts")

    value = 1.0 if not problems else 0.0
    gate = GATES["M6"]
    return MetricResult(
        name="M6_determinism",
        value=value,
        gate=gate.threshold,
        comparison=gate.comparison,
        passed=gate.passes(value),
        detail={"fixtures": list(M6_FIXTURES), "problems": problems},
    )


# --------------------------------------------------------------------------
# M7 — trap disposition
# --------------------------------------------------------------------------


def _candidate_set(item: dict[str, Any]) -> set[str]:
    proposal = item.get("proposal", {})
    if "formats" in proposal:
        return set(proposal["formats"])
    if "readings" in proposal:
        return set(proposal["readings"])
    return set()


def _expectation_satisfied(outcome: Outcome, expectation: dict[str, Any]) -> tuple[bool, str]:
    cells = {(row, col) for row, col in expectation["cells"]}
    klass = expectation["klass"]
    mode = expectation["expect"]

    if expectation.get("expect_type"):
        column = expectation["col_name"]
        actual = outcome.profiles.get(column)
        if actual != expectation["expect_type"]:
            return False, f"column {column} typed {actual}, expected {expectation['expect_type']}"

    if mode == "review":
        rule = expectation["expect_rule"]
        for item in outcome.review_items:
            if item["rule"] != rule:
                continue
            covered = {
                (cell["row"], cell["col"]) for cell in item.get("proposal", {}).get("cells", [])
            }
            if not cells <= covered:
                continue
            wanted = expectation.get("expect_candidates")
            if wanted is not None and _candidate_set(item) != set(wanted):
                return False, f"candidates {sorted(_candidate_set(item))} != {sorted(wanted)}"
            return True, ""
        return False, f"no review item {rule} covering {sorted(cells)}"

    if mode == "report":
        rule = expectation["expect_rule"]
        reported = {
            (finding["row"], finding["col"])
            for finding in outcome.findings
            if finding["rule"] == rule and finding["tier"] == "report"
        }
        missing = cells - reported
        return (not missing), f"missing report-tier {rule} at {sorted(missing)}"

    # mode == "none": no finding of this class at any listed cell, at any tier.
    offenders = [
        (finding["row"], finding["col"])
        for finding in outcome.findings
        if finding["klass"] == klass
        and (
            (finding.get("row"), finding.get("col")) in cells
            or (finding.get("col") is None and any(finding.get("row") == r for r, _ in cells))
        )
    ]
    return (not offenders), f"unexpected {klass} findings at {offenders}"


def m7_trap_disposition(outcomes: list[Outcome]) -> MetricResult:
    by_name = {outcome.name: outcome for outcome in outcomes if outcome.kind == "trap"}
    satisfied = 0
    total = 0
    failures: list[str] = []
    for trap_name, expectations in sorted(forbidden().items()):
        outcome = by_name.get(trap_name)
        for expectation in expectations:
            total += 1
            if outcome is None:
                failures.append(f"{trap_name}: fixture missing")
                continue
            ok, reason = _expectation_satisfied(outcome, expectation)
            if ok:
                satisfied += 1
            else:
                failures.append(f"{trap_name}:{expectation['id']} — {reason}")
    value = satisfied / total if total else 0.0
    gate = GATES["M7"]
    return MetricResult(
        name="M7_trap_disposition",
        value=value,
        gate=gate.threshold,
        comparison=gate.comparison,
        passed=gate.passes(value),
        detail={"expectations": total, "satisfied": satisfied, "failures": failures},
    )


# --------------------------------------------------------------------------
# Report-only metrics (printed, not gated)
# --------------------------------------------------------------------------


def review_precision(outcomes: list[Outcome]) -> dict[str, Any]:
    """Share of review proposals whose recommended candidate equals golden."""
    per_rule: dict[str, list[int]] = defaultdict(lambda: [0, 0])
    for outcome in outcomes:
        if outcome.kind != "corrupted":
            continue
        for item in outcome.review_items:
            for cell in item.get("proposal", {}).get("cells", []):
                row, col = cell.get("row"), cell.get("col")
                if row is None or col is None:
                    continue
                bucket = per_rule[item["rule"]]
                bucket[1] += 1
                if cell.get("after") == outcome.golden_cell(row, col):
                    bucket[0] += 1
    overall_ok = sum(bucket[0] for bucket in per_rule.values())
    overall_n = sum(bucket[1] for bucket in per_rule.values())
    return {
        "overall": round(overall_ok / overall_n, 6) if overall_n else 0.0,
        "per_rule": {
            rule: round(bucket[0] / bucket[1], 6) for rule, bucket in sorted(per_rule.items())
        },
    }


def tier_distribution(outcomes: list[Outcome]) -> dict[str, dict[str, int]]:
    table: dict[str, Counter[str]] = defaultdict(Counter)
    for outcome in outcomes:
        for finding in outcome.findings:
            table[finding["klass"]][finding["tier"]] += 1
    return {klass: dict(sorted(counts.items())) for klass, counts in sorted(table.items())}


def per_file_table(outcomes: list[Outcome]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for outcome in outcomes:
        if outcome.kind != "corrupted":
            continue
        auto_cells, auto_rows = _auto_repairs(outcome)
        review_cells = _review_repairs(outcome)
        ops = injected_ops(outcome)
        fixable = [op for op in ops if op_class(op) not in UNFIXABLE_CLASSES]
        restored = auto_cells | review_cells
        repaired = sum(
            1
            for op in fixable
            if (
                op["row"] in auto_rows
                if op_class(op) == "DUP"
                else (op["row"], op["col"]) in restored
            )
        )
        correct = sum(
            1 for entry in outcome.auto_changes if _auto_change_is_correct(outcome, entry)
        )
        rows.append(
            {
                "fixture": outcome.name,
                "ops": len(ops),
                "auto_changes": len(outcome.auto_changes),
                "auto_precision": round(correct / len(outcome.auto_changes), 4)
                if outcome.auto_changes
                else 1.0,
                "repair_total": round(repaired / len(fixable), 4) if fixable else 0.0,
                "revert": outcome.revert_ok,
            }
        )
    return rows


# --------------------------------------------------------------------------
# The naive baseline (EVALS.md §5) — computed live, never hardcoded
# --------------------------------------------------------------------------


def _naive_read(path: Path) -> Table:
    """Decode UTF-8 with ``errors='replace'`` — the afternoon-script reader."""
    if path.suffix.lower() == ".xlsx":
        return read_table(path)
    text = path.read_bytes().decode("utf-8", errors="replace")
    if path.suffix.lower() == ".jsonl":
        headers: list[str] = []
        records = [
            json.loads(line, parse_float=str, parse_int=str)
            for line in text.splitlines()
            if line.strip()
        ]
        for record in records:
            for key in record:
                if key not in headers:
                    headers.append(key)
        return Table(
            headers=headers,
            rows=[[None if r.get(k) is None else str(r.get(k)) for k in headers] for r in records],
        )
    import csv

    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    grid = [row for row in csv.reader(text.splitlines(), delimiter=delimiter) if row]
    return Table(
        headers=list(grid[0]),
        rows=[[None if cell == "" else cell for cell in row] for row in grid[1:]],
    )


def _naive_type(values: list[str]) -> str:
    """All-or-nothing parse: int → float → %Y-%m-%d → str."""
    if not values:
        return "text"
    try:
        for value in values:
            int(value)
        return "integer"
    except ValueError:
        pass
    try:
        for value in values:
            float(value)
        return "float"
    except ValueError:
        pass
    import datetime as _dt

    try:
        for value in values:
            _dt.date.fromisoformat(value)
        return "date"
    except ValueError:
        return "text"


@dataclass
class NaiveOutcome:
    name: str
    kind: str
    changes: dict[tuple[int, int], str]
    dropped_rows: set[int]
    types: dict[str, str]
    golden_table: Table | None
    manifest: dict[str, Any] | None

    def golden_cell(self, row: int, col: int) -> str | None:
        assert self.golden_table is not None
        return self.golden_table.cell(row, col)


def naive_clean(path: Path, kind: str) -> NaiveOutcome:
    """``str.strip()`` every cell, drop exact duplicate rows, type by parse."""
    table = _naive_read(path)
    changes: dict[tuple[int, int], str] = {}
    stripped: list[list[str | None]] = []
    for row_index, row in enumerate(table.rows):
        new_row: list[str | None] = []
        for col_index, cell in enumerate(row):
            if cell is None:
                new_row.append(None)
                continue
            value = cell.strip()
            if value != cell:
                changes[(row_index, col_index)] = value
            new_row.append(value)
        stripped.append(new_row)

    seen: dict[tuple[str | None, ...], int] = {}
    dropped: set[int] = set()
    for row_index, row in enumerate(stripped):
        signature = tuple(row)
        if signature in seen:
            dropped.add(row_index)
        else:
            seen[signature] = row_index

    types: dict[str, str] = {}
    for col_index, header in enumerate(table.headers):
        values = [
            row[col_index]
            for row_index, row in enumerate(stripped)
            if row_index not in dropped and col_index < len(row) and row[col_index]
        ]
        types[header] = _naive_type([value for value in values if value is not None])

    manifest = manifest_for(path) if kind == "corrupted" else None
    golden = read_table(GOLDEN_DIR / manifest["golden"]) if manifest else read_table(path)
    return NaiveOutcome(
        name=path.name,
        kind=kind,
        changes=changes,
        dropped_rows=dropped,
        types=types,
        golden_table=golden,
        manifest=manifest,
    )


def naive_scores(naive: list[NaiveOutcome]) -> dict[str, float]:
    """Live-computed baseline scores for every gated metric (EVALS.md §5)."""
    labels = type_labels()

    # M1
    correct = total = 0
    for outcome in naive:
        if outcome.kind != "corrupted" or outcome.manifest is None:
            continue
        for column, label in labels[outcome.manifest["golden"]].items():
            total += 1
            correct += int(outcome.types.get(column) == label)
    m1 = correct / total if total else 0.0

    # M2: the naive script "reports" exactly what it changed.
    injected: Counter[str] = Counter()
    detected: Counter[str] = Counter()
    reports: Counter[str] = Counter()
    matched: Counter[str] = Counter()
    for outcome in naive:
        if outcome.kind != "corrupted" or outcome.manifest is None:
            continue
        by_cell = {}
        by_row = {}
        for op in outcome.manifest["ops"]:
            klass = OP_CLASS[op["op"]]
            injected[klass] += 1
            if klass == "DUP":
                by_row[op["row"]] = op
            else:
                by_cell[(op["row"], op["col"])] = op
        for cell in outcome.changes:
            reports["WS"] += 1
            op = by_cell.get(cell)
            if op is not None and OP_CLASS[op["op"]] == "WS":
                matched["WS"] += 1
                detected["WS"] += 1
        for row in outcome.dropped_rows:
            reports["DUP"] += 1
            if row in by_row:
                matched["DUP"] += 1
                detected["DUP"] += 1
    f1s = []
    for klass in M2_CLASSES:
        recall = detected[klass] / injected[klass] if injected[klass] else 0.0
        precision = matched[klass] / reports[klass] if reports[klass] else 0.0
        f1s.append(2 * precision * recall / (precision + recall) if (precision + recall) else 0.0)
    m2 = sum(f1s) / len(f1s)
    m2_min = min(f1s)

    # M3 / M4
    auto_total = auto_correct = 0
    fixable = repaired = 0
    class_total: Counter[str] = Counter()
    class_repaired: Counter[str] = Counter()
    for outcome in naive:
        if outcome.kind != "corrupted" or outcome.manifest is None:
            continue
        duplicates = {op["row"] for op in outcome.manifest["ops"] if op["op"] == "duplicate_row"}
        for (row, col), after in outcome.changes.items():
            auto_total += 1
            auto_correct += int(after == outcome.golden_cell(row, col))
        for row in outcome.dropped_rows:
            auto_total += 1
            auto_correct += int(row in duplicates)
        for op in outcome.manifest["ops"]:
            klass = OP_CLASS[op["op"]]
            if klass in UNFIXABLE_CLASSES:
                continue
            fixable += 1
            class_total[klass] += 1
            if klass == "DUP":
                ok = op["row"] in outcome.dropped_rows
            else:
                key = (op["row"], op["col"])
                ok = key in outcome.changes and outcome.changes[key] == outcome.golden_cell(*key)
            repaired += int(ok)
            class_repaired[klass] += int(ok)
    m3 = auto_correct / auto_total if auto_total else 1.0
    m4 = repaired / fixable if fixable else 0.0
    m4_min = min(
        (class_repaired[klass] / class_total[klass] if class_total[klass] else 0.0)
        for klass in REPAIR_MIN_CLASSES
    )

    # M3_trap / M3_clean / M3_clean_findings
    trap_hits = 0
    for outcome in naive:
        if outcome.kind != "trap":
            continue
        banned = forbidden_cells(outcome.name)
        trap_hits += sum(1 for cell in outcome.changes if cell in banned)
        trap_hits += sum(
            1 for row in outcome.dropped_rows if any(cell[0] == row for cell in banned)
        )
    clean_changes = sum(
        len(outcome.changes) + len(outcome.dropped_rows)
        for outcome in naive
        if outcome.kind == "golden"
    )

    # M7: no review items and no findings, so only `expect: "none"` can pass —
    # the vacuous half of the contract.
    satisfied = total_expectations = 0
    naive_by_name = {outcome.name: outcome for outcome in naive if outcome.kind == "trap"}
    for trap_name, expectations in forbidden().items():
        outcome = naive_by_name.get(trap_name)
        for expectation in expectations:
            total_expectations += 1
            if outcome is None or expectation["expect"] != "none":
                continue
            if (
                expectation.get("expect_type")
                and outcome.types.get(expectation["col_name"]) != expectation["expect_type"]
            ):
                continue
            satisfied += 1

    return {
        "M1": round(m1, 4),
        "M2": round(m2, 4),
        "M2_min": round(m2_min, 4),
        "M3": round(m3, 4),
        "M3_trap": float(trap_hits),
        "M3_clean": float(clean_changes),
        "M3_clean_findings": float(clean_changes),
        "M4_auto": round(m4, 4),
        "M4_auto_min": round(m4_min, 4),
        # It has no review tier at all, so nothing is added to repair_auto.
        "M4_total": round(m4, 4),
        # No audit log at all, so nothing is reconstructible: a run is
        # reversible only if the naive cleaner happened to change nothing.
        "M5": round(
            sum(
                1
                for outcome in naive
                if outcome.kind in {"corrupted", "golden"}
                and not outcome.changes
                and not outcome.dropped_rows
            )
            / max(sum(1 for o in naive if o.kind in {"corrupted", "golden"}), 1),
            4,
        ),
        "M6": 0.0,
        "M7": round(satisfied / total_expectations, 4) if total_expectations else 0.0,
    }


# --------------------------------------------------------------------------
# The suite
# --------------------------------------------------------------------------


@dataclass
class EvalReport:
    metrics: list[MetricResult]
    outcomes: list[Outcome]
    naive: dict[str, float]
    report_only: dict[str, Any]

    @property
    def passed(self) -> bool:
        return all(metric.passed for metric in self.metrics)

    def by_name(self, name: str) -> MetricResult:
        for metric in self.metrics:
            if metric.name.startswith(name):
                return metric
        raise KeyError(name)

    def to_json(self) -> dict[str, Any]:
        return {
            "passed": self.passed,
            "metrics": [
                {
                    "name": metric.name,
                    "value": round(metric.value, 6),
                    "gate": metric.gate,
                    "comparison": metric.comparison,
                    "passed": metric.passed,
                }
                for metric in self.metrics
            ],
            "naive_baseline": self.naive,
            "naive_deltas": {
                metric.name: round(metric.value - self.naive.get(_gate_key(metric.name), 0.0), 6)
                for metric in self.metrics
            },
            "report_only": self.report_only,
        }


def _gate_key(metric_name: str) -> str:
    mapping = {
        "M1_type_inference_accuracy": "M1",
        "M2_detection_macro_f1": "M2",
        "M2_min_class_f1": "M2_min",
        "M3_auto_fix_precision": "M3",
        "M3_trap_auto_changes": "M3_trap",
        "M3_clean_auto_changes": "M3_clean",
        "M3_clean_findings": "M3_clean_findings",
        "M4_repair_auto": "M4_auto",
        "M4_repair_auto_min": "M4_auto_min",
        "M4_repair_total": "M4_total",
        "M5_reversibility": "M5",
        "M6_determinism": "M6",
        "M7_trap_disposition": "M7",
    }
    return mapping[metric_name]


@lru_cache(maxsize=1)
def evaluate(include_m6: bool = True) -> EvalReport:
    """Run every fixture once and compute all metrics.  Cached per process."""
    workdir = Path(tempfile.mkdtemp(prefix="datasweep-evals-"))
    try:
        outcomes: list[Outcome] = []
        for path in corrupted_fixtures():
            outcomes.append(run_fixture(path, "corrupted", workdir))
        for path in golden_fixtures():
            outcomes.append(run_fixture(path, "golden", workdir))
        for path in trap_fixtures():
            outcomes.append(run_fixture(path, "trap", workdir))

        revision = revision_two_reversibility(workdir)
        m2, m2_min = m2_detection(outcomes)
        m4_auto, m4_min, m4_total = m4_repair(outcomes)
        metrics = [
            m1_type_inference(outcomes),
            m2,
            m2_min,
            m3_auto_precision(outcomes),
            m3_trap(outcomes),
            m3_clean(outcomes),
            m3_clean_findings(outcomes),
            m4_auto,
            m4_min,
            m4_total,
            m5_reversibility(outcomes, revision),
        ]
        if include_m6:
            metrics.append(m6_determinism(workdir))
        metrics.append(m7_trap_disposition(outcomes))

        naive = naive_scores(
            [naive_clean(path, "corrupted") for path in corrupted_fixtures()]
            + [naive_clean(path, "golden") for path in golden_fixtures()]
            + [naive_clean(path, "trap") for path in trap_fixtures()]
        )
        report_only = {
            "review_precision": review_precision(outcomes),
            "tier_distribution": tier_distribution(outcomes),
            "per_file_table": per_file_table(outcomes),
        }
        return EvalReport(metrics=metrics, outcomes=outcomes, naive=naive, report_only=report_only)
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


__all__ = [
    "GATES",
    "M2_CLASSES",
    "OP_CLASS",
    "EvalReport",
    "Gate",
    "MetricResult",
    "Outcome",
    "evaluate",
    "expected",
    "forbidden",
    "naive_scores",
    "type_labels",
]
