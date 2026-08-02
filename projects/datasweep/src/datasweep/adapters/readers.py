"""The TableReader port: defensive parsing of CSV/TSV, JSONL and XLSX (FR-3).

Readers produce the *parsed original*: header names plus every cell as a
verbatim string (``str | None``), with the file's shape preserved exactly —
ragged rows and duplicate header names included.  Repairing that shape is the
STR pipeline stage's job, under audit, which is what makes ``revert``
reconstruct the parsed original (D4).

Empty CSV/TSV fields are read as the canonical null ``None`` (D8): the format
cannot distinguish an empty field from an absent one, and the round trip
``None → empty field`` is exact.
"""

from __future__ import annotations

import csv
import io
import json
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Protocol, runtime_checkable

from pydantic import BaseModel, ConfigDict

from ..engine.inference import looks_headerless
from ..engine.models import FileFormat, Issue, IssueClass, Policy, RawTable, Tier
from ..errors import FileTooLargeError, UnsupportedFormatError
from .encoding import DetectedEncoding, EncodingDetector, SimpleEncodingDetector

#: FR-3's candidate delimiter set, in tie-break order.
DELIMITERS: tuple[str, ...] = (",", "\t", ";", "|")
#: How many rows the dialect sniffer and the headerless heuristic look at.
SNIFF_ROWS = 50
HEADER_SAMPLE_ROWS = 200


class FileMeta(BaseModel):
    """What sniffing established about a file (persisted as ``Run.dialect``)."""

    model_config = ConfigDict(frozen=True)

    format: FileFormat
    encoding: str | None = None
    had_bom: bool = False
    delimiter: str | None = None
    quotechar: str | None = None
    has_header: bool = True
    sheet: str | int | None = None

    def dialect(self) -> dict[str, Any]:
        return {
            "delimiter": self.delimiter,
            "quotechar": self.quotechar,
            "has_header": self.has_header,
            "sheet": self.sheet,
        }


class ReadResult(BaseModel):
    """FR-3's ``(RawTable, list[Issue])`` plus the finalized meta.

    ``has_header`` can only be decided once the body has been typed (FR-3's
    headerless rule), so the reader hands back the corrected meta rather than
    making the caller re-derive it for ``Run.dialect``.
    """

    model_config = ConfigDict(frozen=True)

    table: RawTable
    issues: list[Issue] = []
    meta: FileMeta


@runtime_checkable
class TableReader(Protocol):
    extensions: tuple[str, ...]

    def sniff(self, path: str, data: bytes) -> FileMeta:  # pragma: no cover - protocol
        ...

    def read(
        self, data: bytes, meta: FileMeta, policy: Policy
    ) -> ReadResult:  # pragma: no cover - protocol
        ...


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------


def _synthetic_headers(width: int) -> list[str]:
    return [f"col_{index}" for index in range(width)]


def _headerless_issue(width: int) -> Issue:
    return Issue(
        klass=IssueClass.STR,
        rule="fix.synthetic_headers",
        tier=Tier.REVIEW,
        row=0,
        col=None,
        col_name=None,
        value=None,
        evidence={
            "reason": "every row-0 cell matches its column's body type",
            "headers": _synthetic_headers(width),
        },
    )


def _finalize(
    header_row: list[str | None],
    body: list[list[str | None]],
    policy: Policy,
) -> tuple[RawTable, list[Issue], bool]:
    """Decide header-vs-data for row 0 and build the parsed original."""
    headers = [("" if cell is None else cell) for cell in header_row]
    sample = body[:HEADER_SAMPLE_ROWS]
    if headers and looks_headerless(headers, sample, policy):
        width = max([len(headers), *(len(row) for row in body)], default=len(headers))
        rows: list[list[str | None]] = [list(header_row), *body]
        return (
            RawTable(headers=_synthetic_headers(width), rows=rows),
            [_headerless_issue(width)],
            False,
        )
    return RawTable(headers=headers, rows=body), [], True


def _check_limits(data: bytes, rows: int, policy: Policy, path: str) -> None:
    limit = policy.general.max_file_mb * 1024 * 1024
    if len(data) > limit:
        raise FileTooLargeError(
            f"{path} is {len(data)} bytes, over the {policy.general.max_file_mb} MB cap",
            size=len(data),
            limit=limit,
        )
    if rows > policy.general.max_rows:
        raise FileTooLargeError(
            f"{path} has {rows} rows, over the {policy.general.max_rows} row cap",
            rows=rows,
            limit=policy.general.max_rows,
        )


# --------------------------------------------------------------------------
# CSV / TSV
# --------------------------------------------------------------------------


def score_delimiter(text: str, delimiter: str) -> tuple[int, float]:
    """(most common column count, share of rows agreeing) over the first rows."""
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, quotechar='"')
    widths: Counter[int] = Counter()
    for index, row in enumerate(reader):
        if index >= SNIFF_ROWS:
            break
        if row:
            widths[len(row)] += 1
    if not widths:
        return (0, 0.0)
    width, count = max(widths.items(), key=lambda item: (item[1], item[0]))
    return (width if width > 1 else 0, count / sum(widths.values()))


def sniff_delimiter(text: str, default: str) -> str:
    """RFC 4180 fallback: max consistent column count over the first 50 rows.

    ``csv.Sniffer`` gets the first word — it reads quoting and multi-character
    patterns well — but its answer is only accepted when the column-count
    evidence agrees, because Sniffer happily returns a delimiter that splits
    every row differently.
    """
    scores = {delimiter: score_delimiter(text, delimiter) for delimiter in DELIMITERS}
    best = max(
        DELIMITERS,
        key=lambda delimiter: (
            scores[delimiter][0],
            scores[delimiter][1],
            delimiter == default,
            -DELIMITERS.index(delimiter),
        ),
    )
    try:
        sniffed = csv.Sniffer().sniff(text[:8192], delimiters="".join(DELIMITERS)).delimiter
    except (csv.Error, IndexError):
        return best
    if sniffed in scores and scores[sniffed][0] == scores[best][0]:
        return sniffed
    return best


class CsvReader:
    """CSV and TSV, with encoding detection and dialect sniffing (FR-3/FR-4)."""

    extensions = (".csv", ".tsv")

    def __init__(self, detector: EncodingDetector | None = None) -> None:
        self.detector = detector or SimpleEncodingDetector()

    def _decode(self, data: bytes) -> tuple[str, DetectedEncoding]:
        detected = self.detector.detect(data)
        return detected.decode(data), detected

    def sniff(self, path: str, data: bytes) -> FileMeta:
        suffix = Path(path).suffix.lower()
        text, detected = self._decode(data)
        default = "\t" if suffix == ".tsv" else ","
        delimiter = sniff_delimiter(text, default)
        return FileMeta(
            format=FileFormat.TSV if suffix == ".tsv" else FileFormat.CSV,
            encoding=detected.name,
            had_bom=detected.had_bom,
            delimiter=delimiter,
            quotechar='"',
            has_header=True,
        )

    def read(self, data: bytes, meta: FileMeta, policy: Policy) -> ReadResult:
        text, _ = self._decode(data)
        reader = csv.reader(
            io.StringIO(text, newline=""),
            delimiter=meta.delimiter or ",",
            quotechar=meta.quotechar or '"',
        )
        rows: list[list[str | None]] = []
        for raw in reader:
            if not raw:
                continue  # a blank line carries no cells
            rows.append([None if cell == "" else cell for cell in raw])
        _check_limits(data, max(len(rows) - 1, 0), policy, "csv input")
        if not rows:
            return ReadResult(table=RawTable(headers=[], rows=[]), issues=[], meta=meta)
        table, issues, has_header = _finalize(rows[0], rows[1:], policy)
        return ReadResult(
            table=table, issues=issues, meta=meta.model_copy(update={"has_header": has_header})
        )


# --------------------------------------------------------------------------
# JSONL
# --------------------------------------------------------------------------


class JsonlReader:
    """Newline-delimited JSON objects; the key union is the header (FR-3)."""

    extensions = (".jsonl",)

    def sniff(self, path: str, data: bytes) -> FileMeta:
        detected = SimpleEncodingDetector().detect(data)
        return FileMeta(
            format=FileFormat.JSONL,
            encoding=detected.name,
            had_bom=detected.had_bom,
            has_header=True,
        )

    def read(self, data: bytes, meta: FileMeta, policy: Policy) -> ReadResult:
        detected = SimpleEncodingDetector().detect(data)
        text = detected.decode(data)
        keys: list[str] = []
        records: list[tuple[dict[str, Any], str]] = []
        for number, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            try:
                # parse_float/parse_int keep the source's exact numeric text,
                # so no float artifact can enter the table (D3).
                record = json.loads(line, parse_float=str, parse_int=str)
            except json.JSONDecodeError as exc:
                raise UnsupportedFormatError(
                    f"line {number} is not valid JSON: {exc.msg}", line=number
                ) from exc
            if not isinstance(record, dict):
                raise UnsupportedFormatError(
                    f"line {number} is not a JSON object; datasweep reads tabular JSONL",
                    line=number,
                )
            for key in record:
                if key not in keys:
                    keys.append(key)
            records.append((record, line))
        _check_limits(data, len(records), policy, "jsonl input")

        issues: list[Issue] = []
        rows: list[list[str | None]] = []
        for row_index, (record, source) in enumerate(records):
            row: list[str | None] = []
            for col_index, key in enumerate(keys):
                value = record.get(key)
                if isinstance(value, dict | list):
                    # Re-serialize from the untouched parse: the scalar pass
                    # stringifies numbers, which would rewrite nested values.
                    nested = json.loads(source).get(key)
                    row.append(json.dumps(nested, ensure_ascii=False, separators=(",", ":")))
                    issues.append(
                        Issue(
                            klass=IssueClass.STR,
                            rule="detect.nested_value",
                            tier=Tier.REPORT,
                            row=row_index,
                            col=col_index,
                            col_name=key,
                            value=None,
                            evidence={"note": "nested JSON serialized to a string cell"},
                        )
                    )
                elif isinstance(value, bool):
                    row.append("true" if value else "false")
                elif value is None:
                    row.append(None)
                else:
                    row.append(str(value))
            rows.append(row)
        return ReadResult(table=RawTable(headers=keys, rows=rows), issues=issues, meta=meta)


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------


def render_excel_value(value: Any) -> str | None:
    """Excel cell → its displayed string, with no float artifacts (FR-3)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)  # shortest round-trip form, never 0.30000000000000004
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        text = value.isoformat()
        return text[:-9] if text.endswith("T00:00:00") else text
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


class XlsxReader:
    """One sheet per run via openpyxl (D14: a core dependency, not an extra)."""

    extensions = (".xlsx",)

    def sniff(self, path: str, data: bytes) -> FileMeta:
        return FileMeta(format=FileFormat.XLSX, encoding=None, has_header=True, sheet=None)

    def read(self, data: bytes, meta: FileMeta, policy: Policy) -> ReadResult:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            selector = meta.sheet if meta.sheet is not None else policy.general.xlsx_sheet
            sheet = self._select_sheet(workbook, selector)
            grid = [
                [render_excel_value(cell) for cell in row]
                for row in sheet.iter_rows(values_only=True)
            ]
        finally:
            workbook.close()

        while grid and all(cell is None for cell in grid[-1]):
            grid.pop()
        width = 0
        for row in grid:
            trimmed = len(row)
            while trimmed and row[trimmed - 1] is None:
                trimmed -= 1
            width = max(width, trimmed)
        grid = [row[:width] for row in grid]
        _check_limits(data, max(len(grid) - 1, 0), policy, "xlsx input")
        if not grid:
            return ReadResult(table=RawTable(headers=[], rows=[]), issues=[], meta=meta)

        table, issues, has_header = _finalize(grid[0], grid[1:], policy)
        return ReadResult(
            table=table,
            issues=issues,
            meta=meta.model_copy(update={"has_header": has_header, "sheet": sheet.title}),
        )

    @staticmethod
    def _select_sheet(workbook: Any, selector: int | str) -> Any:
        if isinstance(selector, str):
            if selector not in workbook.sheetnames:
                raise UnsupportedFormatError(
                    f"sheet {selector!r} not found; available: {workbook.sheetnames}",
                    sheet=selector,
                )
            return workbook[selector]
        index = int(selector) - 1  # the policy's xlsx_sheet is 1-based
        if not 0 <= index < len(workbook.sheetnames):
            raise UnsupportedFormatError(
                f"sheet index {selector} out of range (1..{len(workbook.sheetnames)})",
                sheet=selector,
            )
        return workbook[workbook.sheetnames[index]]


# --------------------------------------------------------------------------
# registry
# --------------------------------------------------------------------------


def default_readers() -> dict[str, TableReader]:
    csv_reader = CsvReader()
    xlsx_reader = XlsxReader()
    jsonl_reader = JsonlReader()
    return {
        ".csv": csv_reader,
        ".tsv": csv_reader,
        ".xlsx": xlsx_reader,
        ".jsonl": jsonl_reader,
    }


def reader_for(path: str, registry: dict[str, TableReader] | None = None) -> TableReader:
    """Look a reader up by extension; unknown extensions are a clear error."""
    registry = registry if registry is not None else default_readers()
    suffix = Path(path).suffix.lower()
    reader = registry.get(suffix)
    if reader is None:
        raise UnsupportedFormatError(
            f"no reader for {suffix or path!r}; supported: {sorted(registry)}",
            extension=suffix,
        )
    return reader
